"""Integração ponta-a-ponta cruzando as Fases 3-7A (RF02→RF06) — Fase 8 (QA).

Lacuna identificada na auditoria de QA da Fase 8: os testes existentes cobrem
cada camada isoladamente (``test_relatorio_service.py`` já cruza IRRF+match,
``test_controllers.py`` cobre o ``AppController`` ponta-a-ponta mas sem valores
numéricos nem releitura do arquivo exportado). Este módulo fecha essa lacuna
com UM fluxo contínuo, sem inventar comportamento novo:

    ingestão (fake, único ponto de IO substituído) → extração híbrida REAL
    (Fase 4) → cálculo de IRRF REAL (Fase 5) → match de portfólio REAL com
    pendência (Fase 6) → exportação REAL para .xlsx/.pdf em disco (Fase 7A) →
    releitura do .xlsx exportado para provar que o valor sobrevive à
    serialização, não só ao objeto de domínio em memória.

Dados 100% fictícios (fixture sintética já usada em outros testes da suíte).
"""
from __future__ import annotations

from openpyxl import load_workbook

from contract_parser.application.contract_extraction_service import ExtratorContrato
from contract_parser.application.document_ingestor import IngestaoResumo
from contract_parser.domain.documento_texto import DocumentoTexto
from contract_parser.domain.empresa import Empresa
from contract_parser.presentation.controllers import AppController
from tests.support.contract_fixtures import RESIDENCIAL_PF_PJ, carregar_contrato
from tests.support.fakes import FakeEmpresaRepository

# Locatário da fixture RESIDENCIAL_PF_PJ ("Comércio de Alimentos Boa Mesa
# LTDA", CNPJ 11.222.333/0001-81 no texto) — normalizado (14 dígitos).
_CNPJ_LOCATARIO_FIXTURE = "11222333000181"
# Cadastrada no portfólio SEM contrato no lote -> gera pendência real (Fase 6).
_CNPJ_SEM_CONTRATO = "45566778000109"
_RAZAO_SEM_CONTRATO = "Beta Servicos ME"


class _FakeIngestor:
    """Único fake do pipeline: a leitura de arquivos (Fase 3), já testada
    isoladamente em ``tests/application/test_document_ingestor.py``. A partir
    daqui o motor de negócio roda de verdade."""

    def __init__(self, resumo: IngestaoResumo) -> None:
        self._resumo = resumo

    def ingerir(self, pasta):
        return self._resumo


def _app() -> AppController:
    repo = FakeEmpresaRepository()
    repo.add(
        Empresa(
            cnpj=_CNPJ_LOCATARIO_FIXTURE,
            razao_social="Comercio de Alimentos Boa Mesa LTDA",
        )
    )
    repo.add(Empresa(cnpj=_CNPJ_SEM_CONTRATO, razao_social=_RAZAO_SEM_CONTRATO))

    texto = carregar_contrato(RESIDENCIAL_PF_PJ)
    resumo = IngestaoResumo(
        total_arquivos=1,
        processados=1,
        documentos=[DocumentoTexto(caminho="c1.pdf", hash="h1", texto=texto, metodo="nativo")],
    )
    return AppController(repo, ingestor=_FakeIngestor(resumo), extrator=ExtratorContrato())


def test_pipeline_completo_ingestao_a_painel_e_conformidade():
    """Ingestão(fake) -> extração real (F4) -> IRRF real (F5) -> match real (F6)."""
    app = _app()
    app.processar_pasta("qualquer/pasta")

    # Fase 4+5: aluguel R$ 3.500,00, locador PF -> desconto simplificado de
    # R$607,20 (ADR-004) -> base tributável 2.892,80, que cai na faixa 15% /
    # dedução 394,16 da tabela 2026: 2892.80*0.15 - 394.16 = 39,76 (tabela
    # padrão). Rendimento bruto <= R$5.000 -> redutor da Lei nº 15.270/2025
    # (ADR-003) zera o imposto: 39,76 - 312,89 -> max(0, ...) = 0,00.
    linha = app.relatorio.linhas_painel()[0]
    assert linha.locatario == "Comércio de Alimentos Boa Mesa LTDA"
    assert linha.irrf == "R$ 0,00"
    assert linha.revisao is False

    # Fase 6: a segunda empresa cadastrada não tem contrato no lote -> pendência
    # com a mensagem exata do PRD (mesmo formato validado em test_relatorio_service).
    conformidade = app.relatorio.resumo_conformidade()
    assert conformidade.total_empresas_cadastradas == 2
    assert conformidade.total_encontrados == 1
    assert conformidade.pendencias == [
        f"Contrato da Empresa {_RAZAO_SEM_CONTRATO} / {_CNPJ_SEM_CONTRATO} não encontrado"
    ]


def test_pipeline_completo_exportacao_excel_sobrevive_a_releitura(tmp_path):
    """Fase 7A: exporta .xlsx real e relê do disco — não só o objeto em memória."""
    app = _app()
    app.processar_pasta("qualquer/pasta")

    destino = app.relatorio.exportar("excel", tmp_path / "relatorio.xlsx")
    wb = load_workbook(destino)
    try:
        aba_contratos = wb["Contratos"]
        primeira_linha_dados = [c.value for c in aba_contratos[2]]
        assert primeira_linha_dados[1] == _CNPJ_LOCATARIO_FIXTURE  # coluna "CNPJ"
        # Reduzido a R$ 0,00 pelo redutor da Lei nº 15.270/2025 (ver comentário acima).
        assert primeira_linha_dados[4] == "R$ 0,00"  # coluna "IRRF Retido"

        aba_conformidade = wb["Conformidade"]
        textos = [
            str(celula.value)
            for linha in aba_conformidade.iter_rows()
            for celula in linha
            if celula.value is not None
        ]
        assert any(
            f"{_RAZAO_SEM_CONTRATO} / {_CNPJ_SEM_CONTRATO} não encontrado" in t for t in textos
        )
    finally:
        wb.close()


def test_pipeline_completo_exportacao_pdf_nao_quebra(tmp_path):
    """Mesmo pipeline pelo backend PDF (reportlab) — não lança e grava conteúdo."""
    app = _app()
    app.processar_pasta("qualquer/pasta")

    destino = app.relatorio.exportar("pdf", tmp_path / "relatorio.pdf")
    assert destino.exists()
    assert destino.stat().st_size > 500
