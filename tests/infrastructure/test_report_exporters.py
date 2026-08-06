"""Testes dos exportadores de relatório (RF06) — camada infrastructure.

Gera .xlsx/.pdf em ``tmp_path`` e reabre para conferir cabeçalhos/valores.
Também trava a formatação pt-BR (moeda/data). Dados fictícios.
"""
from __future__ import annotations

from datetime import date
from decimal import Decimal

import pytest

from contract_parser.domain.contrato import TipoParte
from contract_parser.domain.irrf import calcular_irrf, tabela_irrf_2026
from contract_parser.domain.relatorio import (
    LinhaContrato,
    Relatorio,
    RelatorioContratos,
    ResumoConformidade,
)
from contract_parser.infrastructure.report_exporters import (
    CABECALHOS_CONTRATOS,
    ExcelRelatorioExporter,
    PdfRelatorioExporter,
    formatar_data_br,
    formatar_moeda_brl,
)

PENDENCIA = "Contrato da Empresa Gamma Holding SA / 00000000000388 não encontrado"


def _relatorio() -> Relatorio:
    tabela = tabela_irrf_2026()
    irrf_pf = calcular_irrf(
        base_mensal=Decimal("5000.00"), tipo_locador=TipoParte.PF, tabela=tabela
    )
    irrf_pj = calcular_irrf(
        base_mensal=Decimal("8000.00"), tipo_locador=TipoParte.PJ, tabela=tabela
    )
    linhas = [
        LinhaContrato(
            locatario_nome="Alpha Comercio LTDA",
            locatario_cnpj="00000000000159",
            locador_nome="João da Silva",
            valor_aluguel=Decimal("5000.00"),
            irrf=irrf_pf,
            indice="IPCA",
            proximo_reajuste="10/2026",
            reajuste_automatico=True,
            vencimento=date(2028, 10, 10),
        ),
        LinhaContrato(
            locatario_nome="Beta Servicos ME",
            locatario_cnpj="00000000000230",
            locador_nome="Imobiliaria XPTO LTDA",
            valor_aluguel=Decimal("8000.00"),
            irrf=irrf_pj,
            indice="IGP-M",
            proximo_reajuste="05/2026",
            reajuste_automatico=False,
            vencimento=date(2027, 5, 1),
        ),
    ]
    return Relatorio(
        contratos=RelatorioContratos(linhas=linhas, tabela_vigencia="2026"),
        conformidade=ResumoConformidade(
            total_empresas_cadastradas=3,
            total_contratos_localizados=2,
            total_encontrados=2,
            pendencias=[PENDENCIA],
        ),
    )


# --------------------------------------------------------------------------- #
# Formatação pt-BR
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize(
    ("valor", "esperado"),
    [
        (Decimal("5000.00"), "R$ 5.000,00"),
        (Decimal("466.27"), "R$ 466,27"),
        (Decimal("1234567.89"), "R$ 1.234.567,89"),
        (Decimal("0.00"), "R$ 0,00"),
        (None, ""),
    ],
)
def test_formatar_moeda_brl(valor, esperado):
    assert formatar_moeda_brl(valor) == esperado


def test_formatar_data_br():
    assert formatar_data_br(date(2028, 10, 10)) == "10/10/2028"
    assert formatar_data_br(None) == ""


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def test_excel_gera_duas_abas_com_cabecalhos_e_valores(tmp_path):
    from openpyxl import load_workbook

    destino = tmp_path / "relatorio.xlsx"
    resultado = ExcelRelatorioExporter().exportar(_relatorio(), destino)

    assert resultado == destino
    assert destino.exists()

    wb = load_workbook(destino)
    assert wb.sheetnames == ["Contratos", "Conformidade"]

    contratos = wb["Contratos"]
    assert [c.value for c in contratos[1]] == CABECALHOS_CONTRATOS
    # Primeira linha de dados: Alpha, valor e IRRF formatados em pt-BR.
    assert contratos["A2"].value == "Alpha Comercio LTDA"
    assert contratos["D2"].value == "R$ 5.000,00"
    # 466,27 (tabela) reduzido a 153,38 pelo redutor da Lei nº 15.270/2025 (ADR-003).
    assert contratos["E2"].value == "R$ 153,38"
    # Coluna nova: valor da redução aplicada (466,27 − 153,38 = 312,89).
    assert contratos["F2"].value == "R$ 312,89"
    assert contratos["I2"].value == "Sim"
    assert contratos["J2"].value == "10/10/2028"
    # Locador PJ ⇒ IRRF R$ 0,00 e nenhuma redução.
    assert contratos["E3"].value == "R$ 0,00"
    assert contratos["F3"].value == "R$ 0,00"

    # Rodapé: total de IRRF retido (153,38 + 0,00) e total de redução (312,89 + 0,00).
    linhas_valores = list(contratos.iter_rows(values_only=True))
    total_row = linhas_valores[-1]
    assert total_row[0] == "TOTAL"
    assert total_row[4] == "R$ 153,38"
    assert total_row[5] == "R$ 312,89"

    conformidade = wb["Conformidade"]
    valores = {row[0].value for row in conformidade.iter_rows() if row[0].value}
    assert "Empresas cadastradas" in valores
    assert PENDENCIA in valores
    wb.close()


def test_excel_cria_diretorio_pai_inexistente(tmp_path):
    destino = tmp_path / "sub" / "dir" / "relatorio.xlsx"
    ExcelRelatorioExporter().exportar(_relatorio(), destino)
    assert destino.exists()


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def test_pdf_gera_arquivo_nao_trivial_com_texto_esperado(tmp_path):
    destino = tmp_path / "relatorio.pdf"
    resultado = PdfRelatorioExporter().exportar(_relatorio(), destino)

    assert resultado == destino
    assert destino.exists()
    # Não-trivial: cabeçalho PDF válido e tamanho razoável.
    dados = destino.read_bytes()
    assert dados[:5] == b"%PDF-"
    assert len(dados) > 1500

    # Texto extraível (via PyMuPDF): confere título, uma linha e a pendência.
    import fitz

    doc = fitz.open(destino)
    try:
        texto = "\n".join(pagina.get_text() for pagina in doc)
    finally:
        doc.close()
    assert "Relatório de Contratos de Locação" in texto
    assert "Alpha Comercio LTDA" in texto
    # 466,27 (tabela) reduzido a 153,38 pelo redutor da Lei nº 15.270/2025 (ADR-003).
    assert "R$ 153,38" in texto
    # Coluna nova: valor da redução aplicada (466,27 − 153,38 = 312,89).
    assert "Redução IRRF" in texto
    assert "R$ 312,89" in texto
    assert PENDENCIA in texto


def test_pdf_sem_pendencias_registra_nenhuma(tmp_path):
    relatorio = Relatorio(
        contratos=RelatorioContratos(linhas=[], tabela_vigencia="2026"),
        conformidade=ResumoConformidade(
            total_empresas_cadastradas=1,
            total_contratos_localizados=1,
            total_encontrados=1,
            pendencias=[],
        ),
    )
    destino = tmp_path / "sem_pendencias.pdf"
    PdfRelatorioExporter().exportar(relatorio, destino)

    import fitz

    doc = fitz.open(destino)
    try:
        texto = "\n".join(pagina.get_text() for pagina in doc)
    finally:
        doc.close()
    assert "Nenhuma pendência." in texto
