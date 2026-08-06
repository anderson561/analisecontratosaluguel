"""Exportadores dos relatórios (camada infrastructure) — RF06.

Materializam o :class:`~contract_parser.domain.relatorio.Relatorio` de domínio
em arquivos, cada um implementando o Protocol ``RelatorioExporter``:

- :class:`ExcelRelatorioExporter` → ``.xlsx`` (openpyxl) com as abas
  "Contratos" (Relatório 01) e "Conformidade" (Relatório 02).
- :class:`PdfRelatorioExporter`   → ``.pdf`` (reportlab) com a tabela de
  contratos e uma seção destacada de pendências.

⚠️ IMPORTS PREGUIÇOSOS (decisão de arquitetura, como ``text_extractors``):
``openpyxl`` e ``reportlab`` são importados DENTRO dos métodos ``exportar`` —
importar este módulo NÃO exige as libs pesadas, e o domínio/aplicação
permanecem independentes do formato de saída.

Formatação pt-BR (R$ 1.234,56 / dd/mm/aaaa) vive aqui, na APRESENTAÇÃO, e é
feita sem depender de ``locale`` (frágil/variável entre SOs).
"""
from __future__ import annotations

from datetime import date
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from contract_parser.domain.relatorio import (
    LinhaContrato,
    Relatorio,
    RelatorioContratos,
    ResumoConformidade,
)

# Cabeçalhos do Relatório 01 (ordem canônica do PRD/§4).
CABECALHOS_CONTRATOS = [
    "Empresa (Locatário)",
    "CNPJ",
    "Locador",
    "Valor Aluguel",
    "IRRF Retido",
    "Redução IRRF (Lei 15.270/2025)",
    "Índice",
    "Próximo Reajuste",
    "Reajuste Auto",
    "Vencimento",
]

_VAZIO = ""  # célula em branco para dado ausente (revisão manual)


# --------------------------------------------------------------------------- #
# Formatação pt-BR (apresentação)
# --------------------------------------------------------------------------- #
def formatar_moeda_brl(valor: Decimal | None) -> str:
    """Formata um ``Decimal`` como moeda pt-BR: ``R$ 1.234,56`` (None → vazio)."""
    if valor is None:
        return _VAZIO
    q = Decimal(valor).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    negativo = q < 0
    inteiro, _, centavos = f"{abs(q):.2f}".partition(".")
    # Agrupa milhares com ponto, da direita para a esquerda.
    partes: list[str] = []
    while len(inteiro) > 3:
        partes.insert(0, inteiro[-3:])
        inteiro = inteiro[:-3]
    partes.insert(0, inteiro)
    milhar = ".".join(partes)
    sinal = "-" if negativo else ""
    return f"{sinal}R$ {milhar},{centavos}"


def formatar_data_br(valor: date | None) -> str:
    """Formata uma data como ``dd/mm/aaaa`` (None → vazio)."""
    if valor is None:
        return _VAZIO
    return valor.strftime("%d/%m/%Y")


def _sim_nao(valor: bool) -> str:
    return "Sim" if valor else "Não"


def _texto(valor: str | None) -> str:
    return valor if valor else _VAZIO


def linha_para_celulas(linha: LinhaContrato) -> list[str]:
    """Converte uma :class:`LinhaContrato` na sequência de células apresentáveis."""
    return [
        _texto(linha.locatario_nome),
        _texto(linha.locatario_cnpj),
        _texto(linha.locador_nome),
        formatar_moeda_brl(linha.valor_aluguel),
        formatar_moeda_brl(linha.irrf_retido),
        formatar_moeda_brl(linha.reducao_irrf),
        _texto(linha.indice),
        _texto(linha.proximo_reajuste),
        _sim_nao(linha.reajuste_automatico),
        formatar_data_br(linha.vencimento),
    ]


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
class ExcelRelatorioExporter:
    """Exporta o relatório para ``.xlsx`` (openpyxl, import lazy)."""

    def _aba_contratos(self, ws, contratos: RelatorioContratos) -> None:
        from openpyxl.styles import Font

        ws.title = "Contratos"
        ws.append(CABECALHOS_CONTRATOS)
        for celula in ws[1]:
            celula.font = Font(bold=True)
        for linha in contratos.linhas:
            ws.append(linha_para_celulas(linha))
        # Rodapé com o total de IRRF retido ("IRRF Retido") e de redução aplicada
        # ("Redução IRRF"), cada total na coluna correspondente ao seu cabeçalho.
        ws.append([])
        ws.append(
            [
                "TOTAL",
                _VAZIO,
                _VAZIO,
                _VAZIO,
                formatar_moeda_brl(contratos.total_irrf_retido),
                formatar_moeda_brl(contratos.total_reducao_irrf),
            ]
        )

    def _aba_conformidade(self, ws, conformidade: ResumoConformidade) -> None:
        from openpyxl.styles import Font

        ws.title = "Conformidade"
        titulo = ws.cell(row=1, column=1, value="Resumo de Conformidade do Portfólio")
        titulo.font = Font(bold=True, size=14)

        totais = [
            ("Empresas cadastradas", conformidade.total_empresas_cadastradas),
            ("Contratos localizados", conformidade.total_contratos_localizados),
            ("Empresas com contrato", conformidade.total_encontrados),
            ("Pendências", conformidade.total_pendencias),
        ]
        linha = 3
        for rotulo, valor in totais:
            ws.cell(row=linha, column=1, value=rotulo).font = Font(bold=True)
            ws.cell(row=linha, column=2, value=valor)
            linha += 1

        linha += 1
        cab = ws.cell(row=linha, column=1, value="Pendências (contratos não encontrados)")
        cab.font = Font(bold=True, size=12)
        linha += 1
        if conformidade.pendencias:
            for pendencia in conformidade.pendencias:
                ws.cell(row=linha, column=1, value=pendencia)
                linha += 1
        else:
            ws.cell(row=linha, column=1, value="Nenhuma pendência.")

    def exportar(self, relatorio: Relatorio, destino: Path) -> Path:
        from openpyxl import Workbook

        destino = Path(destino)
        wb = Workbook()
        try:
            self._aba_contratos(wb.active, relatorio.contratos)
            self._aba_conformidade(wb.create_sheet("Conformidade"), relatorio.conformidade)
            destino.parent.mkdir(parents=True, exist_ok=True)
            wb.save(destino)
        finally:
            wb.close()
        return destino


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
class PdfRelatorioExporter:
    """Exporta o relatório para ``.pdf`` (reportlab, import lazy)."""

    def _tabela_contratos(self, contratos: RelatorioContratos):
        from reportlab.lib import colors
        from reportlab.platypus import Table, TableStyle

        dados = [CABECALHOS_CONTRATOS]
        for linha in contratos.linhas:
            dados.append(linha_para_celulas(linha))
        dados.append(
            [
                "TOTAL",
                "",
                "",
                "",
                formatar_moeda_brl(contratos.total_irrf_retido),
                formatar_moeda_brl(contratos.total_reducao_irrf),
                "",
                "",
                "",
                "",
            ]
        )

        tabela = Table(dados, repeatRows=1)
        tabela.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 7),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F2F2F2")]),
                    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )
        return tabela

    def _bloco_conformidade(self, conformidade: ResumoConformidade):
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, Spacer, Table, TableStyle

        estilos = getSampleStyleSheet()
        flowables = [
            Spacer(1, 18),
            Paragraph("Resumo de Conformidade do Portfólio", estilos["Heading2"]),
        ]

        totais = Table(
            [
                ["Empresas cadastradas", str(conformidade.total_empresas_cadastradas)],
                ["Contratos localizados", str(conformidade.total_contratos_localizados)],
                ["Empresas com contrato", str(conformidade.total_encontrados)],
                ["Pendências", str(conformidade.total_pendencias)],
            ]
        )
        totais.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                ]
            )
        )
        flowables += [Spacer(1, 6), totais, Spacer(1, 12)]

        # Seção destacada de pendências.
        flowables.append(Paragraph("Pendências — contratos não encontrados", estilos["Heading3"]))
        if conformidade.pendencias:
            alerta = estilos["Normal"].clone("alerta")
            alerta.textColor = colors.HexColor("#C00000")
            for pendencia in conformidade.pendencias:
                flowables.append(Paragraph(f"⚠ {pendencia}", alerta))
        else:
            flowables.append(Paragraph("Nenhuma pendência.", estilos["Normal"]))
        return flowables

    def exportar(self, relatorio: Relatorio, destino: Path) -> Path:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        destino = Path(destino)
        destino.parent.mkdir(parents=True, exist_ok=True)

        estilos = getSampleStyleSheet()
        documento = SimpleDocTemplate(
            str(destino),
            pagesize=landscape(A4),
            title="Relatório de Contratos de Locação",
        )
        flowables = [
            Paragraph("Relatório de Contratos de Locação", estilos["Title"]),
            Paragraph(
                f"Relatório 01 — Contratos processados (IRRF tabela {relatorio.contratos.tabela_vigencia})",
                estilos["Heading2"],
            ),
            Spacer(1, 6),
            self._tabela_contratos(relatorio.contratos),
        ]
        flowables += self._bloco_conformidade(relatorio.conformidade)
        documento.build(flowables)
        return destino
