"""Backends de extração de texto (camada infrastructure) — RF02.

Implementa os Protocols de ``domain.extracao`` tocando as bibliotecas pesadas:
- ``PdfExtractor``  → PDF nativo via PyMuPDF (``fitz``), com rota de fallback OCR
  quando a camada de texto é insuficiente (PDF escaneado/imagem).
- ``DocxExtractor`` → DOCX via ``python-docx``.
- ``TesseractOcr``  → renderização de páginas (PyMuPDF) + ``pytesseract`` (pt-BR).

⚠️ IMPORTS PREGUIÇOSOS (decisão travada de arquitetura): ``fitz``, ``pytesseract``,
``PIL`` e ``docx`` são importados DENTRO dos métodos, nunca no topo do módulo.
Motivo: a suíte de testes precisa importar/rodar mesmo sem essas libs ou sem os
binários de SO (Tesseract/poppler) instalados. A seleção do backend e o
roteamento nativo→OCR são testados com fakes injetados.

Resiliência (requisitos §6): arquivo corrompido/ilegível NÃO lança — vira um
``DocumentoTexto`` com ``erro`` preenchido, para que a ingestão em lote continue.
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

from contract_parser.config import settings
from contract_parser.domain.documento_texto import DocumentoTexto, calcular_hash
from contract_parser.domain.extracao import OcrEngine, ResultadoOcr

# Nº mínimo de caracteres "úteis" na camada nativa para considerá-la válida.
# Abaixo disso presume-se PDF escaneado/imagem e cai-se na rota OCR.
LIMIAR_TEXTO_NATIVO = 20

# Confiança média (0–100) do Tesseract abaixo da qual marcamos revisão manual.
LIMIAR_CONFIANCA_OCR = 60.0

# DPI de renderização das páginas para OCR (equilíbrio nitidez × custo).
DPI_OCR = 300


def ler_pdf_nativo(dados: bytes) -> tuple[str, int]:
    """Extrai a camada de texto nativa de um PDF via PyMuPDF (import lazy).

    Retorna ``(texto, n_paginas)``. Levanta a exceção da lib se o PDF for
    inválido — o chamador (``PdfExtractor``) converte isso em ``erro``.
    """
    import fitz  # lazy: só carrega PyMuPDF quando há PDF real para processar

    doc = fitz.open(stream=dados, filetype="pdf")
    try:
        partes = [pagina.get_text() for pagina in doc]
        paginas = doc.page_count
    finally:
        doc.close()
    return "\n".join(partes).strip(), paginas


def _tesseract_embutido() -> tuple[Path, Path] | None:
    """Resolve o Tesseract embutido pelo PyInstaller, se aplicável.

    Só retorna algo quando a aplicação está rodando como executável congelado
    (``sys.frozen`` e ``sys._MEIPASS`` presentes) — o `.spec`
    (``contract_parser.spec``) embute ``tesseract.exe`` + DLLs de runtime +
    ``tessdata/por.traineddata`` em ``<_MEIPASS>/tesseract/``. Fora do modo
    congelado (dev/pytest via ``python -m``) devolve ``None`` — o
    comportamento permanece EXATAMENTE o mesmo de antes desta função existir.
    """
    if not (getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS")):
        return None
    base = Path(sys._MEIPASS) / "tesseract"  # type: ignore[attr-defined]
    return base / "tesseract.exe", base / "tessdata"


class TesseractOcr:
    """Motor de OCR (implementa ``OcrEngine``): PyMuPDF renderiza, Tesseract lê.

    Idioma e caminho do binário vêm de ``settings`` (``ocr_lang`` / ``tesseract_cmd``).
    Todos os imports pesados (``fitz``/``pytesseract``/``PIL``) são feitos dentro de
    ``reconhecer`` — construir o objeto NÃO exige Tesseract instalado.

    Resolução do binário do Tesseract em ``reconhecer`` (ordem de prioridade):
    1. ``self._cmd`` (de ``TESSERACT_CMD``, configurado explicitamente pelo
       usuário via `.env`) — SEMPRE tem prioridade, mesmo rodando como `.exe`
       (permite apontar para um Tesseract do sistema com outros idiomas).
    2. Caso contrário, e SÓ quando rodando como app congelado pelo
       PyInstaller, usa o Tesseract embutido no bundle automaticamente (sem
       exigir configuração do usuário), apontando para o ``tessdata``
       embutido via variável de ambiente ``TESSDATA_PREFIX`` — mecanismo
       oficial do próprio Tesseract, herdado automaticamente pelo subprocesso
       (``pytesseract`` invoca via ``subprocess`` sem ``env=`` explícito, logo
       o processo filho herda ``os.environ`` do processo Python atual).
       NÃO usamos mais ``--tessdata-dir`` via parâmetro ``config``: no
       Windows, ``pytesseract`` monta os argumentos de linha de comando com
       ``shlex.split(config, posix=False)``, e nesse modo as aspas NÃO são
       removidas do token — um caminho entre aspas (necessário para suportar
       espaços) chegava ao Tesseract com as aspas literais embutidas,
       corrompendo a resolução do arquivo de idioma (bug de produção
       reproduzido e confirmado).
    3. Fora do modo congelado, sem ``TESSERACT_CMD``, o comportamento é o
       mesmo de sempre: ``pytesseract`` resolve o binário via ``PATH`` do
       sistema.
    """

    def __init__(
        self,
        *,
        lang: str | None = None,
        tesseract_cmd: str | None = None,
        dpi: int = DPI_OCR,
        limiar_confianca: float = LIMIAR_CONFIANCA_OCR,
    ) -> None:
        self._lang = lang if lang is not None else settings.ocr_lang
        self._cmd = tesseract_cmd if tesseract_cmd is not None else settings.tesseract_cmd
        self._dpi = dpi
        self._limiar = limiar_confianca

    def reconhecer(self, dados: bytes) -> ResultadoOcr:
        import io

        import fitz  # lazy
        import pytesseract  # lazy
        from PIL import Image  # lazy

        if self._cmd:
            # Configuração explícita do usuário (TESSERACT_CMD) tem prioridade
            # sobre o Tesseract embutido, mesmo rodando como `.exe`. Não
            # mexemos em TESSDATA_PREFIX aqui — pode haver um Tesseract de
            # sistema já configurado pelo usuário com seu próprio ambiente.
            pytesseract.pytesseract.tesseract_cmd = self._cmd
        else:
            embutido = _tesseract_embutido()
            if embutido is not None:
                tesseract_exe, tessdata_dir = embutido
                pytesseract.pytesseract.tesseract_cmd = str(tesseract_exe)
                # Variável de ambiente oficial do Tesseract para resolver o
                # diretório do tessdata — o subprocesso lançado por
                # pytesseract herda os.environ automaticamente. Robusto a
                # espaços no caminho, ao contrário de passar via `config`.
                os.environ["TESSDATA_PREFIX"] = str(tessdata_dir)

        textos: list[str] = []
        confiancas: list[float] = []
        doc = fitz.open(stream=dados, filetype="pdf")
        try:
            paginas = doc.page_count
            for pagina in doc:
                pixmap = pagina.get_pixmap(dpi=self._dpi)
                imagem = Image.open(io.BytesIO(pixmap.tobytes("png")))
                dados_ocr = pytesseract.image_to_data(
                    imagem,
                    lang=self._lang,
                    output_type=pytesseract.Output.DICT,
                )
                palavras = [p for p in dados_ocr["text"] if p and p.strip()]
                textos.append(" ".join(palavras))
                confiancas.extend(
                    float(c) for c in dados_ocr["conf"] if str(c) != "-1" and float(c) >= 0
                )
        finally:
            doc.close()

        media = sum(confiancas) / len(confiancas) if confiancas else 0.0
        baixa_confianca = (not confiancas) or media < self._limiar
        return ResultadoOcr(
            texto="\n".join(textos).strip(),
            paginas=paginas,
            baixa_confianca=baixa_confianca,
        )


class PdfExtractor:
    """Extrator de PDF: tenta a camada nativa; cai para OCR se for insuficiente.

    Colaboradores injetáveis (Dependency Injection) tornam o roteamento
    nativo→OCR testável com fakes, sem exigir PyMuPDF/Tesseract:
    - ``leitor_nativo``: ``(bytes) -> (texto, paginas)`` (default: PyMuPDF).
    - ``ocr``: motor ``OcrEngine`` (default: ``TesseractOcr`` — imports lazy só na
      hora de reconhecer, então construir aqui não exige o binário).
    """

    def __init__(
        self,
        *,
        leitor_nativo=ler_pdf_nativo,
        ocr: OcrEngine | None = None,
        limiar_texto: int = LIMIAR_TEXTO_NATIVO,
    ) -> None:
        self._leitor_nativo = leitor_nativo
        self._ocr = ocr if ocr is not None else TesseractOcr()
        self._limiar_texto = limiar_texto

    def aceita(self, caminho: Path) -> bool:
        return caminho.suffix.lower() == ".pdf"

    def extrair(self, caminho: Path) -> DocumentoTexto:
        dados = caminho.read_bytes()
        hash_ = calcular_hash(dados)

        try:
            texto_nativo, paginas = self._leitor_nativo(dados)
        except Exception as exc:  # noqa: BLE001 - PDF corrompido não aborta o lote
            return DocumentoTexto(
                caminho=str(caminho),
                hash=hash_,
                erro=f"Falha ao abrir/ler PDF: {exc!r}",
            )

        if len(texto_nativo) >= self._limiar_texto:
            return DocumentoTexto(
                caminho=str(caminho),
                hash=hash_,
                texto=texto_nativo,
                metodo="nativo",
                paginas=paginas,
            )

        # Camada nativa insuficiente → presume escaneado/imagem: rota OCR.
        try:
            resultado = self._ocr.reconhecer(dados)
        except Exception as exc:  # noqa: BLE001 - falha de OCR não aborta o lote
            return DocumentoTexto(
                caminho=str(caminho),
                hash=hash_,
                paginas=paginas,
                erro=f"Falha no OCR: {exc!r}",
            )
        return DocumentoTexto(
            caminho=str(caminho),
            hash=hash_,
            texto=resultado.texto,
            metodo="ocr",
            paginas=resultado.paginas or paginas,
            baixa_confianca=resultado.baixa_confianca,
        )


class DocxExtractor:
    """Extrator de DOCX via ``python-docx`` (import lazy).

    Concatena parágrafos e o texto das células de tabelas — contratos costumam
    ter quadros de partes/valores em tabelas. ``paginas`` = 0 (DOCX não expõe
    paginação de forma confiável).
    """

    def aceita(self, caminho: Path) -> bool:
        return caminho.suffix.lower() == ".docx"

    def extrair(self, caminho: Path) -> DocumentoTexto:
        from docx import Document  # lazy

        dados = caminho.read_bytes()
        hash_ = calcular_hash(dados)

        try:
            documento = Document(str(caminho))
        except Exception as exc:  # noqa: BLE001 - DOCX corrompido não aborta o lote
            return DocumentoTexto(
                caminho=str(caminho),
                hash=hash_,
                erro=f"Falha ao abrir/ler DOCX: {exc!r}",
            )

        partes = [p.text for p in documento.paragraphs if p.text and p.text.strip()]
        for tabela in documento.tables:
            for linha in tabela.rows:
                celulas = [c.text.strip() for c in linha.cells if c.text and c.text.strip()]
                if celulas:
                    partes.append(" | ".join(celulas))

        return DocumentoTexto(
            caminho=str(caminho),
            hash=hash_,
            texto="\n".join(partes).strip(),
            metodo="docx",
            paginas=0,
        )
