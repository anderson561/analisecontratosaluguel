"""Builders programáticos de PDFs e DOCX para os testes (sem binários versionados).

Gera arquivos determinísticos em ``tmp_path``. Os imports de ``fitz``/``docx`` são
feitos dentro das funções para que o módulo possa ser importado mesmo se as libs
não estiverem instaladas — os testes que chamam estes builders usam
``pytest.importorskip`` antes.
"""
from __future__ import annotations

from pathlib import Path


def construir_pdf_nativo(caminho: Path, texto: str = "CONTRATO DE LOCACAO RESIDENCIAL") -> Path:
    """PDF com camada de texto real (PyMuPDF) — exercita a rota nativa."""
    import fitz

    doc = fitz.open()
    pagina = doc.new_page()
    pagina.insert_text((72, 72), texto)
    doc.save(str(caminho))
    doc.close()
    return caminho


def construir_pdf_sem_texto(caminho: Path, paginas: int = 1) -> Path:
    """PDF válido porém SEM camada de texto (páginas em branco) — força a rota OCR."""
    import fitz

    doc = fitz.open()
    for _ in range(paginas):
        doc.new_page()
    doc.save(str(caminho))
    doc.close()
    return caminho


def construir_docx(caminho: Path, paragrafos: list[str], tabela: list[list[str]] | None = None) -> Path:
    """DOCX real (python-docx) com parágrafos e uma tabela opcional."""
    from docx import Document

    documento = Document()
    for p in paragrafos:
        documento.add_paragraph(p)
    if tabela:
        tab = documento.add_table(rows=len(tabela), cols=len(tabela[0]))
        for i, linha in enumerate(tabela):
            for j, valor in enumerate(linha):
                tab.rows[i].cells[j].text = valor
    documento.save(str(caminho))
    return caminho
